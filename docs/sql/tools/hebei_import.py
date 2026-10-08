# -*- coding: utf-8 -*-
"""河北志愿填报数据入库(用户提供的文件, 见 docs/功能说明.md 第7节)

  1. score_ranking  ← 2021-2026年 河北高考一分一档表.xlsx (物理类/历史类, 替换旧 OCR 数据)
  2. control_line   ← docs/sql/tools/data/hebei_control_lines.csv
  3. recruit_major + recruit_major_history ← 4 个志愿数据 xlsx
       预测最低分 = score_line 当年官方投档最低分; 位次/百分位/等效分由一分一段换算

用法: python hebei_import.py            (全部)
      python hebei_import.py rank line  (只跑部分: rank / line / major)
重复执行幂等(按省份+年份覆盖)。
"""
import bisect, csv, os, re, sys, collections
import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), 'crawler'))
from common import db, ledger_upsert  # noqa: E402

PROVINCE = '河北'
TARGET_YEAR = 2026
SRC_DIR = r'F:\教育考试院\河北省'
RANK_XLSX = os.path.join(SRC_DIR, '一分一档数据', '2021-2026年 河北高考一分一档表.xlsx')
LINE_CSV = os.path.join(os.path.dirname(__file__), 'data', 'hebei_control_lines.csv')
MAJOR_FILES = [
    ('2026年河北本科物理组志愿数据.xlsx', '本科批', '物理类'),
    ('2026年河北本科历史组志愿数据.xlsx', '本科批', '历史类'),
    ('2026年河北高考物理科目组合专科数据含投档线位次.xlsx', '专科批', '物理类'),
    ('2026年河北高考历史科目组合专科数据含投档线位次.xlsx', '专科批', '历史类'),
]
HIST_YEARS = [2025, 2024, 2023, 2022, 2021]

UNV_TYPE_MAP = {'理工': '工科', '航空': '工科', '能源': '工科', '中医药': '医药', '农林': '农业', '文艺': '艺术'}
NATURE_MAP = {'独立学院': '民办', '境外高校海南办学': '中外合作办学'}
SPECIAL_TYPES = {'地方专项计划', '八协计划', '少数民族预科班', '边防军人子女预科班'}


# ---------------------------------------------------------------- 一分一段
def load_rank_tables():
    """{(year, '物理类'|'历史类'): [(score, count, cum)] 分数降序}"""
    wb = openpyxl.load_workbook(RANK_XLSX, read_only=True, data_only=True)
    tables = {}
    for ws in wb.worksheets:
        if not ws.title.isdigit():
            continue
        year = int(ws.title)
        phy, his = [], []
        for r in list(ws.iter_rows(values_only=True))[2:]:
            if r[0] is None:
                continue
            m = re.match(r'\d+', str(r[0]))
            if not m:
                continue
            score = int(m.group())
            if r[2] is not None:
                phy.append((score, int(r[1] or 0), int(r[2])))
            if r[4] is not None:
                his.append((score, int(r[3] or 0), int(r[4])))
        tables[(year, '物理类')] = phy
        tables[(year, '历史类')] = his
    return tables


class Ranker:
    """分数↔位次/百分位/等效分换算"""

    def __init__(self, tables, lines):
        self.tables = tables
        self.lines = lines  # {(year, subject, batch): (score, cumulative)}
        self.cum_by_score = {k: {s: c for s, _, c in t} for k, t in tables.items()}
        self.asc_cums = {k: [c for _, _, c in t] for k, t in tables.items()}

    def rank(self, year, subject, score):
        if score is None:
            return None
        t = self.tables.get((year, subject))
        if not t:
            return None
        m = self.cum_by_score[(year, subject)]
        if score in m:
            return m[score]
        if score > t[0][0]:
            return t[0][2]
        upper = [c for sc, _, c in t if sc >= score]
        return upper[-1] if upper and score >= t[-1][0] else None

    def denominator(self, year, subject, batch):
        v = self.lines.get((year, subject, batch))
        return v[1] if v else None

    def percentile(self, year, subject, batch, rank):
        n = self.denominator(year, subject, batch)
        if rank is None or not n:
            return None
        return round((1 - rank / n) * 100, 4)

    def equivalent(self, year, subject, batch, rank):
        """往年位次 → 目标年百分位最接近的分数档"""
        if rank is None:
            return None
        n_src = self.denominator(year, subject, batch)
        n_dst = self.denominator(TARGET_YEAR, subject, batch)
        t = self.tables.get((TARGET_YEAR, subject))
        if not n_src or not n_dst or not t:
            return None
        target_cum = rank / n_src * n_dst
        cums = self.asc_cums[(TARGET_YEAR, subject)]
        i = bisect.bisect_left(cums, target_cum)
        cands = [j for j in (i - 1, i) if 0 <= j < len(t)]
        best = min(cands, key=lambda j: abs(cums[j] - target_cum))
        return t[best][0]


def import_rank(conn, tables):
    cur = conn.cursor()
    cur.execute('DELETE FROM score_ranking WHERE province=%s', (PROVINCE,))
    data = []
    for (year, subject), rows in sorted(tables.items()):
        for s, c, cum in rows:
            data.append((PROVINCE, year, subject, s, c, cum))
    cur.executemany('INSERT INTO score_ranking (province,year,subject_type,score,same_count,cumulative) '
                    'VALUES (%s,%s,%s,%s,%s,%s)', data)
    conn.commit()
    for (year, subject), rows in sorted(tables.items()):
        ledger_upsert(conn, PROVINCE, year, 'score_ranking', '', subject, file_url=os.path.basename(RANK_XLSX),
                      local_file=RANK_XLSX, file_format='xlsx', rows_loaded=len(rows), status='loaded',
                      note='用户提供(含优惠分)')
    print(f'score_ranking: {len(data)} 行 ({len(tables)} 组)')


def load_lines():
    out = {}
    with open(LINE_CSV, encoding='utf-8') as f:
        for r in csv.DictReader(f):
            out[(int(r['year']), r['subject_type'], r['batch'])] = (int(r['score']), int(r['rank']), r['source'])
    return out


def import_lines(conn, lines):
    cur = conn.cursor()
    cur.execute('DELETE FROM control_line WHERE province=%s', (PROVINCE,))
    cur.executemany('INSERT INTO control_line (province,year,subject_type,batch,score,cumulative,source) '
                    'VALUES (%s,%s,%s,%s,%s,%s,%s)',
                    [(PROVINCE, y, s, b, v[0], v[1], v[2]) for (y, s, b), v in sorted(lines.items())])
    conn.commit()
    print(f'control_line: {len(lines)} 行')


# ---------------------------------------------------------------- 专业代码映射
class MajorMapper:
    def __init__(self, conn):
        cur = conn.cursor()
        cur.execute('SELECT major_code, major_name, first_discipline FROM major')
        self.by_name = {}
        self.by_disc = collections.defaultdict(list)
        for code, name, disc in cur.fetchall():
            self.by_name.setdefault(name, code)
            self.by_disc[disc].append(code)
        self.miss = 0

    @staticmethod
    def base(name):
        return re.sub(r'[（(][^（()）]*[)）]\s*$', '', name or '').strip()

    def note_names(self, note):
        m = re.search(r'含(.+?)(?:等)?专业', note or '')
        if not m:
            return []
        return [x.strip() for x in re.split(r'[、，,/]', m.group(1)) if x.strip()]

    def codes(self, name, note):
        b = self.base(name)
        if name in self.by_name:
            return [self.by_name[name]]
        if b in self.by_name:
            return [self.by_name[b]]
        out = []
        for n in self.note_names(note):
            c = self.by_name.get(n) or self.by_name.get(self.base(n))
            if c and c not in out:
                out.append(c)
        if out:
            return out
        if b.endswith('类') and b in self.by_disc:
            return list(self.by_disc[b])
        self.miss += 1
        return []


# ---------------------------------------------------------------- 志愿数据
def s(v):
    if v is None:
        return ''
    return str(v).strip()


def n(v):
    """数值列: 空/0/非数字 → None"""
    if v is None:
        return None
    try:
        f = float(str(v).strip().replace(',', ''))
    except ValueError:
        return None
    return int(f) if f else None


def unv_attributes(level):
    tags = []
    if '双一流' in level:
        tags.append('双一流')
    if '/985' in level:
        tags.append('985')
    if '/211' in level:
        tags.append('211')
    return tags


def import_majors(conn, ranker):
    cur = conn.cursor()
    cur.execute("SELECT year, batch, subject_type, college_code, major_code, college_name, min_mark "
                "FROM score_line WHERE province=%s AND year=%s", (PROVINCE, TARGET_YEAR))
    official = {}
    official_unv = {}
    for y, b, st, cc, mc, cn, mm in cur.fetchall():
        k = (b, st, str(cc).zfill(4))
        official[k + (str(mc).zfill(2),)] = mm
        official_unv.setdefault(k, cn)

    mapper = MajorMapper(conn)
    loaded = []
    for fname, batch, subject in MAJOR_FILES:
        path = os.path.join(SRC_DIR, fname)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
        h = rows[0]
        ix = {k: i for i, k in enumerate(h)}
        loaded.append((fname, batch, subject, ix, rows[1:]))

    # 院校级硕博点: 同校任一专业行有即算有
    master_unv, doctor_unv = set(), set()
    for fname, batch, subject, ix, rows in loaded:
        for r in rows:
            code = s(r[ix['院校代码']]).zfill(4)
            if n(r[ix['硕士点']]):
                master_unv.add(code)
            if n(r[ix['博士点']]):
                doctor_unv.add(code)

    yy = str(TARGET_YEAR)[2:]
    recs, hists = [], []
    stat = collections.Counter()
    for fname, batch, subject, ix, rows in loaded:
        g = lambda r, col: r[ix[col]] if col in ix else None  # noqa: E731
        prefix = f"{yy}冀{'本' if batch == '本科批' else '专'}{'理' if subject == '物理类' else '文'}"
        for r in rows:
            unv_code = s(g(r, '院校代码')).zfill(4)
            mcode = s(g(r, '专业代码')).zfill(2)
            uab = prefix + unv_code
            rid = f'{uab}-{mcode}'
            nature_raw = re.sub(r'\(有本科\)', '', s(g(r, '办学性质')))
            nature = NATURE_MAP.get(nature_raw, nature_raw)
            short = s(g(r, '院校名称'))
            unv_name = official_unv.get((batch, subject, unv_code)) or f'{short}[{nature_raw}]'
            level = s(g(r, '院校水平'))
            rtype = s(g(r, '招生类型')) or '普通计划'
            note = s(g(r, '专业简注'))
            label = rtype if rtype in SPECIAL_TYPES else '普通类'
            explain = f'({label})' + (f'({note})' if note else '')
            tuition_raw = g(r, '学费')
            tuition = int(tuition_raw) if isinstance(tuition_raw, (int, float)) else (0 if s(tuition_raw) == '免费' else None)
            second = s(g(r, '次选科目'))
            req = [s(g(r, '首选科目'))] + ([x for x in second.split('&') if x] if second and second != '不限' else [])
            newflag = s(g(r, '新招专业'))
            mname = s(g(r, '专业名称'))
            codes = mapper.codes(mname, note)

            pre_min = official.get((batch, subject, unv_code, mcode))
            pre_rank = ranker.rank(TARGET_YEAR, subject, pre_min)
            pre_pct = ranker.percentile(TARGET_YEAR, subject, batch, pre_rank) if pre_min is not None else None

            has_hist = False
            for y in HIST_YEARS:
                k = str(y)[2:]
                plan, mn, rk = n(g(r, f'{k}人')), n(g(r, f'{k}分')), n(g(r, f'{k}位'))
                avg = n(g(r, f'{k}均分'))
                if plan is None and mn is None and rk is None:
                    continue
                if rk is None and mn is not None:
                    rk = ranker.rank(y, subject, mn)
                avg_rank = ranker.rank(y, subject, avg)
                if mn is not None and y >= TARGET_YEAR - 3:
                    has_hist = True
                hists.append((rid, y, TARGET_YEAR, plan, mn, rk,
                              ranker.percentile(y, subject, batch, rk) if mn is not None else None,
                              ranker.equivalent(y, subject, batch, rk) if mn is not None else None,
                              avg, avg_rank,
                              ranker.percentile(y, subject, batch, avg_rank) if avg is not None else None,
                              ranker.equivalent(y, subject, batch, avg_rank) if avg is not None else None))

            unv_rank = n(g(r, '院校排名'))
            mrank = n(g(r, '软科专业排名')) if '软科专业排名' in ix else n(g(r, '专业排名'))
            stat[(batch, subject)] += 1
            stat['official'] += pre_min is not None
            recs.append((
                rid, uab, PROVINCE, TARGET_YEAR, batch, subject,
                unv_code, unv_name, short, nature, UNV_TYPE_MAP.get(s(g(r, '院校类型')), s(g(r, '院校类型'))),
                s(g(r, '省份')), s(g(r, '城市')), s(g(r, '城市评级')), unv_rank, s(g(r, '软科校排')),
                s(g(r, '主管部门')), level, s(g(r, '院校背景')), ','.join(unv_attributes(level)),
                1 if unv_code in master_unv else 0, 1 if unv_code in doctor_unv else 0,
                1 if '保研资格' in level else 0,
                s(g(r, '校园VR')), s(g(r, '院校百科')), s(g(r, '招生章程')), s(g(r, '学校招生信息')), s(g(r, '就业质量')),
                1 if newflag == '新招院校' else 0,
                mcode, mname, note, s(g(r, '专业类')).lstrip('/'), s(g(r, '注意事项')), rtype, explain,
                n(g(r, '26人')), n(g(r, '学制')), tuition, s(tuition_raw) if tuition is None or tuition == 0 else '',
                ','.join(req), ','.join(codes), s(g(r, '评估')), mrank, s(g(r, '专业水平')),
                n(g(r, '硕士点')) or 0, s(g(r, '硕士专业')), n(g(r, '博士点')) or 0, s(g(r, '博士专业')),
                s(g(r, '26推免')), (float(g(r, '25热度')) if s(g(r, '25热度')) else None),
                1 if newflag == '新招专业' else 0,
                pre_min, pre_rank, pre_pct, n(g(r, '26参考分')), n(g(r, '26参考位')),
                'No' if has_hist else 'Danger',
                n(g(r, '25校低分')), n(g(r, '25校低位')), n(g(r, '25校均分')), n(g(r, '25校均位')),
                fname,
            ))

    cols = ('recruit_major_id,uab_id,province,year,batch,subject_type,unv_code,unv_name,unv_short_name,unv_nature,'
            'unv_type,unv_province,unv_city,city_level,unv_rank,rk_rank,authority,unv_level,unv_background,'
            'unv_attributes,has_master,has_doctor,can_recommend,vr_url,baike_url,charter_url,admission_url,'
            'employment_url,first_time_admission,recruit_major_code,major_name,major_note,major_category,attention,'
            'recruit_type,plan_explain,major_plan,school_length,tuition,tuition_text,required_subjects,major_codes,'
            'major_grade,major_rank,major_level,master_count,master_majors,doctor_count,doctor_majors,tuimian_rate,'
            'heat,is_new_major,pre_min_mark,pre_min_rank,pre_min_percentile,ref_mark,ref_rank,attention_type,'
            'unv_min_mark_last,unv_min_rank_last,unv_avg_mark_last,unv_avg_rank_last,source_file')
    ncol = len(cols.split(','))
    assert all(len(x) == ncol for x in recs), (ncol, len(recs[0]))

    cur.execute('DELETE h FROM recruit_major_history h JOIN recruit_major m ON m.recruit_major_id=h.recruit_major_id '
                'WHERE m.province=%s AND m.year=%s', (PROVINCE, TARGET_YEAR))
    cur.execute('DELETE FROM recruit_major WHERE province=%s AND year=%s', (PROVINCE, TARGET_YEAR))
    sql = f"INSERT INTO recruit_major ({cols}) VALUES ({','.join(['%s'] * ncol)})"
    for i in range(0, len(recs), 2000):
        cur.executemany(sql, recs[i:i + 2000])
    hsql = ('INSERT INTO recruit_major_history (recruit_major_id,year,equivalent_year,plan,min_mark,min_mark_rank,'
            'min_mark_percentile,equivalent_min_mark,avg_mark,avg_mark_rank,avg_mark_percentile,equivalent_avg_mark) '
            'VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)')
    for i in range(0, len(hists), 5000):
        cur.executemany(hsql, hists[i:i + 5000])
    # 格学口径: 首次招生 = 该院校(批次+科类)下所有专业都没有任何历年投档分
    cur.execute('UPDATE recruit_major m LEFT JOIN ('
                ' SELECT x.uab_id FROM recruit_major x JOIN recruit_major_history h'
                ' ON h.recruit_major_id=x.recruit_major_id AND h.min_mark IS NOT NULL'
                ' WHERE x.province=%s AND x.year=%s GROUP BY x.uab_id) t ON t.uab_id=m.uab_id'
                ' SET m.first_time_admission = IF(t.uab_id IS NULL, 1, 0)'
                ' WHERE m.province=%s AND m.year=%s', (PROVINCE, TARGET_YEAR, PROVINCE, TARGET_YEAR))
    conn.commit()

    for fname, batch, subject, ix, rows in loaded:
        ledger_upsert(conn, PROVINCE, TARGET_YEAR, 'plan', batch, subject, file_url=fname,
                      local_file=os.path.join(SRC_DIR, fname), file_format='xlsx', rows_loaded=len(rows),
                      status='loaded', note='用户提供: 招生计划+历年专业分/位次+院校扩展信息')
    print(f'recruit_major: {len(recs)} 行, 其中有当年官方投档分 {stat["official"]} 行')
    for k, v in stat.items():
        if isinstance(k, tuple):
            print('   ', k, v)
    print(f'recruit_major_history: {len(hists)} 行')
    print(f'专业代码未映射: {mapper.miss} 行(专科专业不在教育部本科目录内属正常)')


def main():
    parts = set(sys.argv[1:]) or {'rank', 'line', 'major'}
    conn = db()
    tables = load_rank_tables()
    lines = load_lines()
    if 'rank' in parts:
        import_rank(conn, tables)
    if 'line' in parts:
        import_lines(conn, lines)
    if 'major' in parts:
        import_majors(conn, Ranker(tables, {k: (v[0], v[1]) for k, v in lines.items()}))
    conn.close()


if __name__ == '__main__':
    main()
