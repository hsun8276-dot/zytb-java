# -*- coding: utf-8 -*-
"""河北省教育考试院 本科批平行志愿投档情况统计 xlsx → zytb.score_line
权威来源: hebeea.edu.cn 普通高考栏目 → 平行志愿投档情况统计 → file.hebeea.edu.cn 附件xlsx
文件结构: 第2行标注科类(物理/历史科目组合), 第6行起为数据行
列: 院校代号|院校名称|专业代号|专业名称|投档最低分|语数之和|语数最高|外语|首选科目|再选最高|再选次高|志愿号|备注
用法: python import_hebei_score_line.py <xlsx路径...>
"""
import sys, os, re
import openpyxl
import pymysql

def norm_subject(s):
    if not s: return None
    if "物理" in s: return "物理类"
    if "历史" in s: return "历史类"
    return None

def parse_file(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    # 年份: 第1行标题"20XX年河北省..."
    year_tag = None
    for r in rows[:2]:
        m = re.search(r"(\d{4})年", str(r[0] or ""))
        if m: year_tag = int(m.group(1)); break
    if not year_tag:
        print(f"  [跳过] 无法识别年份: {os.path.basename(path)}")
        return None
    # 科类: 第2行(索引1)出现"物理科目组合"/"历史科目组合"
    subject = None
    for r in rows[:4]:
        joined = "".join(str(v or "") for v in r)
        s = norm_subject(joined)
        if s: subject = s; break
    if not subject:
        print(f"  [跳过] 无法识别科类: {os.path.basename(path)}")
        return None
    # 批次: 标题含"专科批"→专科批, 含"本科批"→本科批; 对口/单招/提前批不要
    title_head = " ".join(str(r[0] or "") for r in rows[:2])
    if "对口" in title_head or "单招" in title_head:
        print(f"  [跳过] 对口/单招数据: {os.path.basename(path)}"); return None
    if "专科提前批" in title_head or "本科提前批" in title_head:
        print(f"  [跳过] 提前批: {os.path.basename(path)}"); return None
    batch = "专科批" if "专科批" in title_head else ("本科批" if "本科批" in title_head else None)
    if not batch:
        print(f"  [跳过] 无法识别批次: {os.path.basename(path)}"); return None
    out = []
    for r in rows:
        code = str(r[0] or "").strip()
        major_code = str(r[2] or "").strip()
        name = str(r[1] or "").strip()
        mname = str(r[3] or "").strip()
        if not code or not major_code or not name or not mname: continue
        try: min_mark = int(float(r[4]))
        except (TypeError, ValueError): continue
        remark = str(r[12] or "").strip() if len(r) > 12 else ""
        out.append((code, name, major_code, mname, min_mark, remark))
    print(f"  {os.path.basename(path)} [{year_tag}/{batch}/{subject}] {len(out)} 行")
    return year_tag, batch, subject, out

def main():
    files = sys.argv[1:]
    if not files:
        print("用法: python import_hebei_score_line.py <xlsx...>"); return
    conn = pymysql.connect(host="127.0.0.1", port=3306, user="root",
                           password="root", database="zytb", charset="utf8mb4")
    cur = conn.cursor()
    for path in files:
        got = parse_file(path)
        if not got: continue
        year_tag, batch, subject, rows = got
        cur.execute("DELETE FROM score_line WHERE province=%s AND year=%s AND batch=%s AND subject_type=%s",
                    ("河北", year_tag, batch, subject))
        ins = []
        for code, cname, mcode, mname, mark, remark in rows:
            # 从院校名提取性质 [公办]/[民办]/[独立学院]
            prop = ""
            m = re.search(r"\[([^\]]+)\]", cname)
            if m: prop = m.group(1)
            try: cc = int(code)
            except ValueError: cc = None
            ins.append(("河北", year_tag, batch, subject, cc, cname, mcode, mname, prop, mark, remark))
        for ofs in range(0, len(ins), 1000):
            cur.executemany("""INSERT INTO score_line (province,year,batch,subject_type,college_code,
                               college_name,major_code,major_name,property,min_mark,remark)
                               VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""", ins[ofs:ofs+1000])
        conn.commit()
        print(f"  → 入库 {year_tag}/{batch}/{subject}: {len(ins)} 行")
    cur.execute("""SELECT year, batch, subject_type, COUNT(*), MIN(min_mark), MAX(min_mark) FROM score_line
                   WHERE province='河北' GROUP BY year, batch, subject_type ORDER BY year, batch, subject_type""")
    print("库内汇总:")
    for row in cur.fetchall(): print("  ", row)
    conn.close()

if __name__ == "__main__":
    main()
