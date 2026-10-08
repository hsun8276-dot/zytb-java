import re
import sys
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

PATH = r"F:\教育考试院\河北省\一分一档数据\2021-2026年 河北高考一分一档表.xlsx"

wb = openpyxl.load_workbook(PATH, data_only=True)
print("工作表:", wb.sheetnames)
for ws in wb.worksheets:
    print("\n" + "=" * 90)
    print(f"[{ws.title}] 维度={ws.dimensions} 最大行={ws.max_row} 最大列={ws.max_column} 状态={ws.sheet_state}")
    if ws.merged_cells.ranges:
        print("  合并单元格:", [str(r) for r in ws.merged_cells.ranges][:30])
    hidden_cols = [k for k, d in ws.column_dimensions.items() if d.hidden]
    hidden_rows = [k for k, d in ws.row_dimensions.items() if d.hidden]
    if hidden_cols or hidden_rows:
        print("  隐藏列:", hidden_cols, "隐藏行:", hidden_rows[:20])
    comments = [(c.coordinate, c.comment.text) for row in ws.iter_rows() for c in row if c.comment]
    if comments:
        print("  批注:", comments)

    if not ws.title.isdigit():
        for row in ws.iter_rows():
            cells = [(c.coordinate, c.value) for c in row if c.value is not None]
            if cells:
                print("  ", cells)
        continue

    used_cols = sorted({c.column_letter for row in ws.iter_rows() for c in row if c.value is not None})
    print("  有值的列:", used_cols)
    rows = list(ws.iter_rows(values_only=True))
    print("  表头:", rows[0][:8], rows[1][:8])
    phy = his = 0
    prev = {}
    problems = []
    for i, r in enumerate(rows[2:], start=3):
        if all(v is None for v in r):
            continue
        extra = [v for v in r[5:] if v is not None]
        if extra:
            problems.append((i, "额外列", extra))
        m = re.match(r"(\d+)", str(r[0]))
        if not m:
            problems.append((i, "非分数行", r[:5]))
            continue
        score = int(m.group(1))
        for name, cnt_i, cum_i in (("物理", 1, 2), ("历史", 3, 4)):
            cnt, cum = r[cnt_i], r[cum_i]
            if cum is None:
                continue
            if name == "物理":
                phy += 1
            else:
                his += 1
            p = prev.get(name)
            if p and cnt is not None and p[1] + cnt != cum:
                problems.append((i, f"{name}累计不连续", score, p, cnt, cum))
            prev[name] = (score, cum)
    print(f"  物理档数={phy} 历史档数={his}")
    print("  首档:", rows[2][:5], "历史首档:", next((r[:5] for r in rows[2:] if r[4] is not None), None))
    print("  末档:", next(r[:5] for r in reversed(rows) if r[0] is not None))
    print(f"  问题 {len(problems)} 条:", problems[:15])
